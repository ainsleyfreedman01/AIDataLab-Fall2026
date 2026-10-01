"""
Builds a color-coded Excel workbook from k12_ai_policies.csv, since plain CSV
files cannot store cell colors/formatting - only a real spreadsheet format
(.xlsx) can. Each state gets one distinct pastel fill color across both its
state-level and district-level rows.

Usage: python3 build_colored_workbook.py
Output: k12_ai_policies.xlsx
"""

import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

csv.field_size_limit(sys.maxsize)

STATE_COLORS = {
    "California": "FDE9D9",
    "Washington": "DCE6F1",
    "Texas": "E2EFDA",
    "New York": "FCE4D6",
    "Pennsylvania": "E4DFEC",
    "Massachusetts": "FFF2CC",
    "Georgia": "D9E1F2",
    "Mississippi": "EDEDED",
    "Alabama": "F8CBAD",
    "West Virginia": "C6E0B4",
    "Kentucky": "BDD7EE",
    "Oklahoma": "FFE699",
    "Wyoming": "D5C4E0",
    "Louisiana": "B4E5E0",
}

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def main() -> None:
    with open("k12_ai_policies.csv", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames
        rows = list(reader)

    wb = Workbook()
    ws = wb.active
    ws.title = "K-12 AI Policies"

    ws.append(fieldnames)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append([row[field] for field in fieldnames])
        fill_color = STATE_COLORS.get(row["state"])
        if fill_color:
            fill = PatternFill("solid", fgColor=fill_color)
            for cell in ws[ws.max_row]:
                cell.fill = fill

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    widths = {"state": 14, "comparison_group": 22, "level": 10, "doc_type": 20, "title": 45, "source_url": 40,
              "date_retrieved": 14, "status": 10, "char_count": 11, "full_text": 60}
    for idx, field in enumerate(fieldnames, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(field, 20)

    ws.freeze_panes = "A2"

    out_path = "k12_ai_policies.xlsx"
    wb.save(out_path)
    print(f"Wrote {len(rows)} color-coded rows to {out_path}")


if __name__ == "__main__":
    main()
